"""9종 돌봄로봇 및 기기 실증 계획 — 표 하나를 앱에서 입력·저장하고 엑셀로 내려받는다.

칸 구성은 정지수 연구원이 자료실에 올린 엑셀의 '26.10.14' 시트와 같다(20칸):
기본 4칸 + 2025년 평가 결과 6칸 + 2026년 평가 계획 10칸.
데이터: 제출함 스프레드시트 `실증계획` 탭(공통확인사항과 같은 '통째로 덮어쓰기' 저장).

⚠️ 담당자 여럿이 같은 표를 고친다 → 화면은 편집 시작 시점의 사본(snapshot)으로 그리고,
저장 직전에 시트를 다시 읽어 그 사이 다른 사람이 저장했는지(rows_hash) 확인한다.
"""
import hashlib
import io
import json

import gspread
import streamlit as st

from sheets_store import _get_client

DEMO_WS = "실증계획"
BASE = ["연번", "분야", "사용성평가 담당자", "세부과제 주관기관"]
RESULT_25 = ["진행기관", "제품 명칭", "제품 버전", "점수", "대상자", "n수"]
PLAN_26 = ["진행기관", "제품 명칭", "제품 버전", "대상자", "n수", "평가 진행 일정",
           "시제품 재활원 입고 예정일", "획득하는 데이터", "평가 완료일", "비고"]
# 시트·화면용 칸 이름(같은 이름이 25/26 양쪽에 있어 앞에 연도를 붙인다)
COLS = BASE + [f"25 {c}" for c in RESULT_25] + [f"26 {c}" for c in PLAN_26]
TITLE = "9종 돌봄로봇 및 기기 실증 계획 (2025년 결과 및 2026년 계획)"
_WIDTHS = [6, 13, 10, 22, 14, 18.9, 10, 8, 12, 7,
           14, 14, 13, 9.4, 12, 12.9, 11.9, 26, 26, 40.5]   # 원본 엑셀 열 너비


@st.cache_resource
def _ws():
    ss = _get_client().open_by_key(st.secrets["sheet"]["id"])
    try:
        ws = ss.worksheet(DEMO_WS)
    except gspread.WorksheetNotFound:
        ws = ss.add_worksheet(title=DEMO_WS, rows=100, cols=len(COLS))
        ws.update(values=[COLS], range_name="A1")
        return ws
    if ws.col_count < len(COLS):
        ws.add_cols(len(COLS) - ws.col_count)
    if ws.row_values(1) != COLS:
        end = gspread.utils.rowcol_to_a1(1, len(COLS))
        ws.update(values=[COLS], range_name=f"A1:{end}")
    return ws


def _clean(vals) -> list:
    out = []
    for r in vals:
        r = [str(c) for c in (list(r) + [""] * len(COLS))[:len(COLS)]]
        if any(c.strip() for c in r):
            out.append(r)
    return out


def load_plan_fresh() -> list:
    """캐시 없이 시트에서 바로 읽는다(저장 전 충돌 확인용)."""
    return _clean(_ws().get_all_values()[1:])


def rows_hash(rows) -> str:
    return hashlib.sha1(json.dumps(_clean(rows), ensure_ascii=False)
                        .encode("utf-8")).hexdigest()


def save_plan(rows) -> int:
    """표 전체를 시트에 덮어쓴다. 저장한 행 수 반환.

    ⚠️ RAW 로 쓴다 — USER_ENTERED 면 '5/5'(n수)·'10월 16일' 같은 글이 날짜로 바뀐다.
    """
    rows = _clean(rows)
    ws = _ws()
    n = len(ws.get_all_values())
    if n > 1:
        ws.batch_clear([f"A2:{gspread.utils.rowcol_to_a1(n, len(COLS))}"])
    if rows:
        if ws.row_count < len(rows) + 1:
            ws.add_rows(len(rows) + 1 - ws.row_count)
        ws.update(values=rows, range_name="A2", value_input_option="RAW")
    return len(rows)


def build_plan_xlsx(rows, as_of: str = "") -> bytes:
    """원본 엑셀(26.10.14 시트)과 같은 모양: 제목 · 2단 머리글(2025 결과/2026 계획) · 표."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = as_of or "실증계획"
    n = len(COLS)
    last = get_column_letter(n)
    font = Font(name="맑은 고딕", size=10)
    bold = Font(name="맑은 고딕", size=10, bold=True)
    head = PatternFill("solid", fgColor="D9E2F3")
    plan = PatternFill("solid", fgColor="FFFF00")    # 원본도 2026 머리글 아래칸이 노랑
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    mid = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A1"] = TITLE + (f" ({as_of} 기준)" if as_of else "")
    ws["A1"].font = Font(name="맑은 고딕", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(f"A1:{last}1")

    e25, e26 = len(BASE) + 1, len(BASE) + len(RESULT_25) + 1
    for i, name in enumerate(BASE, 1):
        ws.cell(3, i, name.replace("사용성평가 ", "사용성평가\n"))
        ws.merge_cells(start_row=3, start_column=i, end_row=4, end_column=i)
    ws.cell(3, e25, "2025년 평가 결과")
    ws.merge_cells(start_row=3, start_column=e25, end_row=3, end_column=e26 - 1)
    ws.cell(3, e26, "2026년 평가 계획")
    ws.merge_cells(start_row=3, start_column=e26, end_row=3, end_column=n)
    for i, name in enumerate(RESULT_25, e25):
        ws.cell(4, i, name)
    for i, name in enumerate(PLAN_26, e26):
        ws.cell(4, i, name.replace("재활원 입고", "재활원\n입고"))
    for r in (3, 4):
        for c in range(1, n + 1):
            cell = ws.cell(r, c)
            cell.font, cell.alignment, cell.border = bold, mid, box
            cell.fill = plan if (r == 4 and c >= e26) else head
    ws.row_dimensions[4].height = 27

    for ri, row in enumerate(_clean(rows), 5):
        for ci, v in enumerate(row, 1):
            cell = ws.cell(ri, ci, _num(v))
            cell.font, cell.alignment, cell.border = font, mid, box
        # 진행기관 칸에만 글이 있으면(예: "해당 없음…") 그 묶음을 가로로 합친다 — 원본과 같게
        for s, e in ((e25, e26 - 1), (e26, n)):
            if row[s - 1].strip() and not any(x.strip() for x in row[s:e]):
                ws.merge_cells(start_row=ri, start_column=s, end_row=ri, end_column=e)
        # 줄 수 = 직접 바꾼 줄 + 열 너비를 넘어 저절로 접히는 줄(한글 1자 ≈ 너비 2)
        lines = max(sum(max(1, -(-sum(2 if ord(ch) > 0x2E7F else 1.1 for ch in ln)
                                   // max(_WIDTHS[ci] - 1.5, 1)))
                        for ln in str(v).split("\n"))
                    for ci, v in enumerate(row))
        ws.row_dimensions[ri].height = max(27, 13.5 * lines + 6)

    for i, w in enumerate(_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(v: str):
    """숫자만 든 칸은 숫자로(점수·연번 정렬·계산되게). 그 외는 글 그대로."""
    s = str(v).strip()
    try:
        return int(s) if s.isdigit() else float(s) if s.replace(".", "", 1).isdigit() else v
    except ValueError:
        return v
